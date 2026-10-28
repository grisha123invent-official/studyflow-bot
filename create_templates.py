"""
Скрипт для создания базовых шаблонов .docx в папке /templates/.
Запускается один раз: python create_templates.py

Создаёт ГОСТ-оформление:
- Times New Roman 14pt
- Поля: лево 30мм, верх 20мм, низ 20мм, право 15мм
- Межстрочный интервал 1.5
"""
import os
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_LINE_SPACING


TEMPLATES_DIR = "templates"


def set_page_margins(doc: Document) -> None:
    """Устанавливает ГОСТ-поля для всего документа"""
    section = doc.sections[0]
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)


def set_paragraph_style(paragraph, font_name: str = "Times New Roman", font_size: int = 14) -> None:
    """Применяет ГОСТ-шрифт и межстрочный интервал к параграфу"""
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = Pt(font_size)

    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)


def add_styled_heading(doc: Document, text: str, level: int = 1) -> None:
    """Добавляет заголовок с нужным форматированием"""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)


def add_styled_paragraph(doc: Document, text: str, first_line_indent: bool = True) -> None:
    """Добавляет параграф с ГОСТ-оформлением"""
    para = doc.add_paragraph(text)
    if first_line_indent:
        para.paragraph_format.first_line_indent = Mm(12.5)  # 1.25 см — стандартный отступ
    set_paragraph_style(para)


def create_base_template(doc: Document, work_type_title: str, structure_items: list[str]) -> None:
    """Наполняет документ базовой структурой для заданного типа работы"""
    set_page_margins(doc)

    # Титульный лист-заглушка
    title = doc.add_paragraph()
    title.alignment = 1  # WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(work_type_title)
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.font.bold = True
    title.paragraph_format.space_before = Pt(72)  # отступ сверху

    # Поле для темы
    doc.add_paragraph()
    field_para = doc.add_paragraph()
    field_para.alignment = 1
    r = field_para.add_run('Тема: "___________________________________"')
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    doc.add_paragraph()
    r2 = doc.add_paragraph()
    r2.alignment = 1
    rr = r2.add_run("Студент: ________________________")
    rr.font.name = "Times New Roman"
    rr.font.size = Pt(14)

    # Разрыв страницы перед содержанием
    doc.add_page_break()

    # Содержание
    add_styled_heading(doc, "СОДЕРЖАНИЕ", level=1)
    for i, item in enumerate(structure_items, 1):
        para = doc.add_paragraph(f"{i}. {item}…………………………… с. ___")
        set_paragraph_style(para)

    doc.add_page_break()

    # Разделы-заглушки
    for item in structure_items:
        add_styled_heading(doc, item.upper(), level=1)
        add_styled_paragraph(doc, "[Здесь будет текст раздела]")
        doc.add_paragraph()


def make_kursovaya() -> None:
    doc = Document()
    structure = [
        "Введение",
        "Глава 1. Теоретическая часть",
        "1.1. [Параграф 1]",
        "1.2. [Параграф 2]",
        "Глава 2. Практическая часть",
        "2.1. [Параграф 1]",
        "2.2. [Параграф 2]",
        "Заключение",
        "Список использованных источников",
    ]
    create_base_template(doc, "КУРСОВАЯ РАБОТА", structure)
    path = os.path.join(TEMPLATES_DIR, "tmpl_kursovaya.docx")
    doc.save(path)
    print(f"✅ {path}")


def make_referat() -> None:
    doc = Document()
    structure = [
        "Введение",
        "Раздел 1. [Название]",
        "Раздел 2. [Название]",
        "Раздел 3. [Название]",
        "Заключение",
        "Список литературы",
    ]
    create_base_template(doc, "РЕФЕРАТ", structure)
    path = os.path.join(TEMPLATES_DIR, "tmpl_referat.docx")
    doc.save(path)
    print(f"✅ {path}")


def make_statya() -> None:
    doc = Document()
    structure = [
        "Аннотация",
        "Ключевые слова",
        "Введение",
        "Основная часть",
        "Выводы",
        "Список литературы",
    ]
    create_base_template(doc, "НАУЧНАЯ СТАТЬЯ", structure)
    path = os.path.join(TEMPLATES_DIR, "tmpl_statya.docx")
    doc.save(path)
    print(f"✅ {path}")


def make_base() -> None:
    doc = Document()
    structure = [
        "Введение",
        "Основная часть",
        "Заключение",
        "Список литературы",
    ]
    create_base_template(doc, "УЧЕБНАЯ РАБОТА", structure)
    path = os.path.join(TEMPLATES_DIR, "tmpl_base.docx")
    doc.save(path)
    print(f"✅ {path}")


if __name__ == "__main__":
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    print(f"Создаём шаблоны в папке {TEMPLATES_DIR}/...\n")
    make_kursovaya()
    make_referat()
    make_statya()
    make_base()
    print("\n🎉 Готово! Все шаблоны созданы.")
