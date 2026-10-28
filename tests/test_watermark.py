"""
Юнит-тесты для services/watermark.py
Запуск: python3 -m pytest tests/test_watermark.py -v
"""
import os
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document


def _make_test_docx(path: str) -> None:
    """Создаёт минимальный .docx с кириллическим текстом для тестов."""
    doc = Document()
    doc.add_heading("Тестовый заголовок", level=1)
    doc.add_paragraph("Это тестовый абзац на русском языке.")
    doc.add_paragraph("Second paragraph in English.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Ячейка 1"
    table.rows[0].cells[1].text = "Ячейка 2"
    table.rows[1].cells[0].text = "Cell 3"
    table.rows[1].cells[1].text = "Cell 4"
    doc.save(path)


class TestPurePythonConversion(unittest.TestCase):
    """Тесты pure-python fallback: python-docx → fpdf2"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.docx_path = os.path.join(self.tmpdir, "test.docx")
        _make_test_docx(self.docx_path)

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_pure_python_creates_pdf(self):
        """pure-python конвертация создаёт PDF-файл"""
        from services.watermark import _convert_docx_to_pdf_pure_python
        pdf_path = os.path.join(self.tmpdir, "test.pdf")
        _convert_docx_to_pdf_pure_python(self.docx_path, pdf_path)
        self.assertTrue(os.path.isfile(pdf_path), "PDF файл должен быть создан")
        self.assertGreater(os.path.getsize(pdf_path), 100, "PDF не должен быть пустым")

    def test_pure_python_pdf_is_valid(self):
        """Созданный PDF открывается через PyMuPDF"""
        import fitz
        from services.watermark import _convert_docx_to_pdf_pure_python
        pdf_path = os.path.join(self.tmpdir, "test.pdf")
        _convert_docx_to_pdf_pure_python(self.docx_path, pdf_path)
        doc = fitz.open(pdf_path)
        self.assertGreater(len(doc), 0, "PDF должен содержать хотя бы 1 страницу")
        doc.close()


class TestConvertChain(unittest.TestCase):
    """Тест цепочки fallback: _convert_docx_to_pdf"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.docx_path = os.path.join(self.tmpdir, "chain_test.docx")
        _make_test_docx(self.docx_path)

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_convert_chain_produces_pdf(self):
        """Цепочка libreoffice→docx2pdf→pure-python создаёт PDF"""
        from services.watermark import _convert_docx_to_pdf
        pdf_path = _convert_docx_to_pdf(self.docx_path, self.tmpdir)
        self.assertTrue(os.path.isfile(pdf_path))
        self.assertGreater(os.path.getsize(pdf_path), 100)

    def test_convert_chain_returns_pdf_extension(self):
        """Цепочка возвращает путь с расширением .pdf"""
        from services.watermark import _convert_docx_to_pdf
        pdf_path = _convert_docx_to_pdf(self.docx_path, self.tmpdir)
        self.assertTrue(pdf_path.endswith(".pdf"))


class TestWatermark(unittest.TestCase):
    """Тесты наложения водяного знака"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.docx_path = os.path.join(self.tmpdir, "wm_test.docx")
        _make_test_docx(self.docx_path)

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_add_watermark_creates_preview(self):
        """add_watermark создаёт _preview.pdf файл"""
        from services.watermark import add_watermark
        result = add_watermark(self.docx_path, order_id=42)
        self.assertTrue(os.path.isfile(result))
        self.assertIn("preview", result)

    def test_add_watermark_pdf_has_pages(self):
        """PDF с водяным знаком содержит страницы"""
        import fitz
        from services.watermark import add_watermark
        result = add_watermark(self.docx_path, order_id=99)
        doc = fitz.open(result)
        self.assertGreater(len(doc), 0)
        doc.close()

    def test_add_watermark_is_flattened(self):
        """Превью растеризовано: нет текстового слоя (знак нельзя удалить) и есть картинка"""
        import fitz
        from services.watermark import add_watermark
        result = add_watermark(self.docx_path, order_id=7)
        doc = fitz.open(result)
        page_text = doc[0].get_text().strip()
        images = doc[0].get_images()
        doc.close()
        self.assertEqual(page_text, "", "В превью не должно быть текстового слоя")
        self.assertGreater(len(images), 0, "Страница должна быть картинкой (растеризована)")


class TestMergePreview(unittest.TestCase):
    """Тесты склейки нескольких файлов в одно превью (add_watermark_merged)."""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _make_pdf(self, name: str, widths: list[int]) -> str:
        import fitz
        path = os.path.join(self.tmpdir, name)
        doc = fitz.open()
        for w in widths:
            doc.new_page(width=w, height=600)
        doc.save(path)
        doc.close()
        return path

    def test_merge_preserves_page_order(self):
        """_merge_pdfs склеивает страницы строго в переданном порядке."""
        import fitz
        from services.watermark import _merge_pdfs
        a = self._make_pdf("a.pdf", [210, 220])
        b = self._make_pdf("b.pdf", [330, 340, 350])
        out = os.path.join(self.tmpdir, "merged.pdf")
        _merge_pdfs([a, b], out)
        doc = fitz.open(out)
        widths = [round(p.rect.width) for p in doc]
        doc.close()
        self.assertEqual(widths, [210, 220, 330, 340, 350],
                         "Порядок страниц при склейке должен сохраняться")

    def test_merged_combines_all_files_and_flattens(self):
        """add_watermark_merged: страницы всех файлов в одном PDF, без текстового слоя."""
        import fitz
        from services.watermark import add_watermark, add_watermark_merged
        f1 = os.path.join(self.tmpdir, "doc1.docx")
        f2 = os.path.join(self.tmpdir, "doc2.docx")
        _make_test_docx(f1)
        _make_test_docx(f2)

        def pages(p):
            d = fitz.open(p); n = len(d); d.close(); return n

        n1 = pages(add_watermark(f1, order_id=11))
        n2 = pages(add_watermark(f2, order_id=11))

        merged = add_watermark_merged([f1, f2], order_id=11, variant=0)
        doc = fitz.open(merged)
        total = len(doc)
        first_text = doc[0].get_text().strip()
        first_imgs = doc[0].get_images()
        doc.close()

        self.assertEqual(total, n1 + n2, "В склейке должны быть страницы всех файлов")
        self.assertEqual(first_text, "", "Склейка тоже растеризована — без текстового слоя")
        self.assertGreater(len(first_imgs), 0, "Страница склейки — картинка")

    def test_merged_single_file_returns_one_preview(self):
        """Один файл — склеивать нечего, возвращается одно превью."""
        from services.watermark import add_watermark_merged
        f1 = os.path.join(self.tmpdir, "solo.docx")
        _make_test_docx(f1)
        result = add_watermark_merged([f1], order_id=12, variant=0)
        self.assertTrue(os.path.isfile(result))
        self.assertTrue(result.endswith(".pdf"))

    def test_merged_empty_raises(self):
        """Пустой список файлов — ошибка, а не молчаливый пустой PDF."""
        from services.watermark import add_watermark_merged
        with self.assertRaises(ValueError):
            add_watermark_merged([], order_id=13)

    def test_pdf_input_watermarks_and_preserves_original(self):
        """PDF на входе: превью с водяным знаком создаётся, оригинал НЕ удаляется."""
        import fitz
        from services.watermark import add_watermark
        src = self._make_pdf("final.pdf", [300, 300])
        result = add_watermark(src, order_id=21, variant=0)
        self.assertTrue(os.path.isfile(result), "Превью должно создаться")
        self.assertTrue(os.path.isfile(src), "Оригинальный PDF не должен удаляться")
        self.assertNotEqual(os.path.abspath(result), os.path.abspath(src))
        doc = fitz.open(result)
        n, txt, imgs = len(doc), doc[0].get_text().strip(), doc[0].get_images()
        doc.close()
        self.assertEqual(n, 2, "Число страниц сохраняется")
        self.assertEqual(txt, "", "Превью растеризовано (знак нельзя удалить)")
        self.assertGreater(len(imgs), 0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
